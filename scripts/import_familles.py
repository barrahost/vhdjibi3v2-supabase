"""
import_familles.py
=================
Lit le fichier "Repartitions Familles V0.1.xlsx" et compare avec les
âmes déjà présentes dans Supabase.

Usage
-----
  Mode analyse (rapport uniquement, rien n'est modifié) :
    python scripts/import_familles.py

  Mode import (insère les âmes absentes après confirmation) :
    python scripts/import_familles.py --import

  Exporter le rapport en Excel :
    python scripts/import_familles.py --excel

"""

import sys
import re
import json
import secrets
import string
import unicodedata
import argparse
from pathlib import Path
from datetime import datetime

import pandas as pd
import requests

# ─── Config ────────────────────────────────────────────────────────────────────

XLSX_PATH = Path(__file__).parent.parent / "Repartitions Familles V0.1.xlsx"
ENV_PATH  = Path(__file__).parent.parent / ".env"

def load_env():
    env = {}
    with open(ENV_PATH) as f:
        for line in f:
            line = line.strip()
            if "=" in line and not line.startswith("#"):
                k, v = line.split("=", 1)
                env[k.strip()] = v.strip().strip('"').strip("'")
    return env

ENV = load_env()
SUPABASE_URL = ENV["VITE_SUPABASE_URL"]
SUPABASE_KEY = ENV["VITE_SUPABASE_PUBLISHABLE_KEY"]
CHURCH_ID    = "bergerie"

_NANOID_ALPHABET = string.ascii_letters + string.digits

def nanoid(size=21):
    return ''.join(secrets.choice(_NANOID_ALPHABET) for _ in range(size))

HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
}

# ─── Helpers ───────────────────────────────────────────────────────────────────

def normalize_phone(p):
    """
    Normalise vers une clé de comparaison : les 8 derniers chiffres ivoiriens.
    Ex: '0161080922' → '61080922', '+2250161080922' → '61080922'
    """
    if pd.isna(p) or str(p).strip() in ("", "nan"):
        return None
    s = re.sub(r"[^\d]", "", str(p))
    # Retirer le préfixe pays 225 si présent
    if s.startswith("225") and len(s) >= 11:
        s = s[3:]
    # Retirer le 0 local initial si présent
    if s.startswith("0") and len(s) == 9:
        s = s[1:]
    # Clé = 8 derniers chiffres
    key = s[-8:] if len(s) >= 8 else s
    return key if len(key) >= 8 else None

def normalize_name(n):
    """Minuscules, sans accents, sans espaces multiples."""
    if not n or str(n) == "nan":
        return ""
    n = str(n).strip()
    n = unicodedata.normalize("NFD", n)
    n = "".join(c for c in n if unicodedata.category(c) != "Mn")
    return " ".join(n.lower().split())

def clean_str(v):
    if pd.isna(v) or str(v).strip() in ("nan", "NaT", ""):
        return None
    return str(v).strip()

def clean_bool(v):
    if pd.isna(v) or str(v).strip() in ("nan", ""):
        return False
    return str(v).strip().lower() in ("oui", "yes", "true", "1")

def clean_date(v):
    if pd.isna(v) or str(v).strip() in ("nan", ""):
        return None
    try:
        d = pd.to_datetime(v, dayfirst=True)
        return d.strftime("%Y-%m-%d")
    except Exception:
        return None

# ─── Lecture du fichier Excel ──────────────────────────────────────────────────

FAMILY_MAP = {
    "Famille CHARIS":       ("CHARIS",        "TUEHI REGINA"),
    "FAMILLE ARBRE DE VIE": ("ARBRE DE VIE",  "ASSOMA JOCELYNE"),
    "FAMILLE SHAMA":        ("SHAMA",          "KPEYA FRED"),
    "FAMILLE CEPHAS":       ("CEPHAS",         "KOUASSI SIMON"),
}

def read_xlsx():
    sheets = pd.read_excel(XLSX_PATH, sheet_name=None, header=None)
    records = []
    for sheet_name, df in sheets.items():
        family_name, family_berger = FAMILY_MAP.get(sheet_name, (sheet_name, ""))
        data = df.iloc[4:].copy()
        data.columns = ["nom", "surnom", "genre", "tel", "lieu",
                        "prem_visite", "berger", "ndn", "baptise",
                        "acad_vdh", "ecole_pdv", "depts"]
        for _, row in data.iterrows():
            nom = clean_str(row["nom"])
            if not nom:
                continue
            tel   = normalize_phone(row["tel"])
            berger_raw = clean_str(row["berger"])
            # Lignes sans date ni téléphone = berger/serviteur, pas âme à importer
            is_shepherd_line = (tel is None and clean_date(row["prem_visite"]) is None
                                and berger_raw is None)
            records.append({
                "famille":     family_name,
                "fam_berger":  family_berger,
                "nom":         nom,
                "nom_norm":    normalize_name(nom),
                "surnom":      clean_str(row["surnom"]),
                "genre":       clean_str(row["genre"]) or "F",
                "tel":         tel,
                "tel_raw":     clean_str(row["tel"]),
                "lieu":        clean_str(row["lieu"]),
                "prem_visite": clean_date(row["prem_visite"]),
                "berger":      berger_raw,
                "ndn":         clean_bool(row["ndn"]),
                "baptise":     clean_bool(row["baptise"]),
                "acad_vdh":    clean_bool(row["acad_vdh"]),
                "ecole_pdv":   clean_bool(row["ecole_pdv"]),
                "depts":       clean_str(row["depts"]),
                "is_shepherd_line": is_shepherd_line,
            })
    return records

# ─── Lecture Supabase ──────────────────────────────────────────────────────────

def fetch_souls():
    url = f"{SUPABASE_URL}/rest/v1/souls"
    params = {
        "select": "id,full_name,phone,spiritual_profile",
        "limit": 2000,
    }
    r = requests.get(url, headers=HEADERS, params=params)
    if r.status_code != 200:
        raise Exception(f"{r.status_code} {r.text[:200]}")
    return r.json()

# ─── Comparaison ───────────────────────────────────────────────────────────────

def compare(records, db_souls):
    by_phone = {}
    by_name  = {}
    for s in db_souls:
        p = normalize_phone(s.get("phone", ""))
        if p:
            by_phone[p] = s
        n = normalize_name(s.get("full_name", ""))
        if n:
            by_name[n] = s

    results = []
    for rec in records:
        if rec["is_shepherd_line"]:
            rec["statut"] = "BERGER/SERVITEUR"
            rec["db_id"] = None
            rec["diff"] = []
            results.append(rec)
            continue

        matched = None
        match_type = None

        if rec["tel"] and rec["tel"] in by_phone:
            matched = by_phone[rec["tel"]]
            match_type = "tel"
        elif rec["nom_norm"] in by_name:
            matched = by_name[rec["nom_norm"]]
            match_type = "nom"

        if matched:
            sp = matched.get("spiritual_profile") or {}
            diffs = []
            if rec["ndn"] and not sp.get("isBornAgain"):
                diffs.append("Né de nouveau: xlsx=Oui / app=Non")
            if rec["baptise"] and not sp.get("isBaptized"):
                diffs.append("Baptisé: xlsx=Oui / app=Non")
            if rec["acad_vdh"] and not sp.get("isEnrolledInAcademy"):
                diffs.append("Académie VDH: xlsx=Oui / app=Non")
            rec["statut"] = "PRÉSENT" if not diffs else "PRÉSENT (diff)"
            rec["db_id"] = matched["id"]
            rec["match_type"] = match_type
            rec["diff"] = diffs
        elif not rec["tel"]:
            rec["statut"] = "SANS TEL (non vérifiable)"
            rec["db_id"] = None
            rec["diff"] = []
        else:
            rec["statut"] = "ABSENT"
            rec["db_id"] = None
            rec["diff"] = []

        results.append(rec)
    return results

# ─── Rapport texte ─────────────────────────────────────────────────────────────

def print_report(results):
    totals = {"PRÉSENT": 0, "PRÉSENT (diff)": 0, "ABSENT": 0,
              "SANS TEL (non vérifiable)": 0, "BERGER/SERVITEUR": 0}

    print("\n" + "="*70)
    print("RAPPORT D'ANALYSE — Repartitions Familles V0.1.xlsx")
    print("="*70)

    current_family = None
    for r in results:
        if r["famille"] != current_family:
            current_family = r["famille"]
            print(f"\n{'─'*60}")
            print(f"  FAMILLE {current_family}  (Berger: {r['fam_berger']})")
            print(f"{'─'*60}")

        s = r["statut"]
        totals[s] = totals.get(s, 0) + 1

        icon = {
            "PRÉSENT":               "✅",
            "PRÉSENT (diff)":        "⚠️ ",
            "ABSENT":                "❌",
            "SANS TEL (non vérifiable)": "❓",
            "BERGER/SERVITEUR":      "👤",
        }.get(s, "  ")

        tel_str = r.get("tel_raw") or "-"
        berger_str = r.get("berger") or ""
        print(f"  {icon}  {r['nom']:<35} {tel_str:<14} [{berger_str[:25]}]")
        if r["diff"]:
            for d in r["diff"]:
                print(f"        ↳ {d}")

    print("\n" + "="*70)
    print("RÉSUMÉ")
    print("="*70)
    total_souls = sum(v for k, v in totals.items() if k != "BERGER/SERVITEUR")
    print(f"  Âmes dans le fichier    : {total_souls}")
    print(f"  ✅ Déjà dans l'app      : {totals['PRÉSENT']}")
    print(f"  ⚠️  Présentes avec diff   : {totals['PRÉSENT (diff)']}")
    print(f"  ❌ Absentes (à importer) : {totals['ABSENT']}")
    print(f"  ❓ Sans tel (non vérifié): {totals['SANS TEL (non vérifiable)']}")
    print(f"  👤 Bergers/Serviteurs   : {totals['BERGER/SERVITEUR']}")
    print()
    return totals

# ─── Export Excel du rapport ───────────────────────────────────────────────────

def export_excel(results):
    rows = []
    for r in results:
        rows.append({
            "Famille":        r["famille"],
            "Statut":         r["statut"],
            "Nom":            r["nom"],
            "Téléphone":      r["tel_raw"] or "",
            "Genre":          r["genre"],
            "Lieu":           r["lieu"] or "",
            "1ère visite":    r["prem_visite"] or "",
            "Berger (xlsx)":  r["berger"] or "",
            "Né de nouveau":  "Oui" if r["ndn"] else "",
            "Baptisé":        "Oui" if r["baptise"] else "",
            "Académie VDH":   "Oui" if r["acad_vdh"] else "",
            "École PDV":      "Oui" if r["ecole_pdv"] else "",
            "Départements":   r["depts"] or "",
            "Différences":    " | ".join(r["diff"]),
            "ID Supabase":    r["db_id"] or "",
        })

    df = pd.DataFrame(rows)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out = Path(__file__).parent.parent / f"rapport_import_{ts}.xlsx"

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Rapport", index=False)
        ws = writer.sheets["Rapport"]

        from openpyxl.styles import Font, PatternFill, Alignment
        # Header
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color="00665C")
            cell.alignment = Alignment(horizontal="center")

        COLOR_MAP = {
            "PRÉSENT":               "D4EDDA",
            "PRÉSENT (diff)":        "FFF3CD",
            "ABSENT":                "F8D7DA",
            "SANS TEL (non vérifiable)": "E2E3E5",
            "BERGER/SERVITEUR":      "D1ECF1",
        }
        for row in ws.iter_rows(min_row=2):
            statut = row[1].value
            color  = COLOR_MAP.get(statut, "FFFFFF")
            for cell in row:
                cell.fill = PatternFill("solid", start_color=color)

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                len(str(col[0].value or "")),
                max((len(str(c.value or "")) for c in col[1:]), default=0)
            ) + 3

    print(f"\n📊 Rapport Excel exporté : {out.name}")
    return out

# ─── Import des âmes absentes ──────────────────────────────────────────────────

def do_import(results):
    to_import = [r for r in results if r["statut"] == "ABSENT" and r["tel"]]
    if not to_import:
        print("Aucune âme à importer (toutes présentes ou sans téléphone).")
        return

    print(f"\n{len(to_import)} âmes à importer. Aperçu des 5 premières :")
    for r in to_import[:5]:
        print(f"  - {r['nom']} | {r['tel_raw']} | {r['famille']}")

    print(f"\nTaper OUI pour confirmer l'import de {len(to_import)} âmes :")
    confirm = input("> ").strip()
    if confirm.upper() != "OUI":
        print("Import annulé.")
        return

    url = f"{SUPABASE_URL}/rest/v1/souls"
    inserted = 0
    errors   = []

    for r in to_import:
        # Formater le téléphone en +225XXXXXXXX
        raw_digits = re.sub(r"[^\d]", "", r["tel_raw"] or "")
        if raw_digits.startswith("225"):
            phone_fmt = "+" + raw_digits
        elif raw_digits.startswith("0") and len(raw_digits) >= 9:
            phone_fmt = "+225" + raw_digits[1:]
        elif len(raw_digits) == 8:
            phone_fmt = "+225" + raw_digits
        else:
            phone_fmt = "+225" + raw_digits

        # Profil spirituel en JSONB
        spiritual_profile = {
            "isBornAgain":             r["ndn"],
            "isBaptized":              r["baptise"],
            "isEnrolledInAcademy":     r["acad_vdh"],
            "isEnrolledInLifeBearers": r["ecole_pdv"],
            "departments":             [],
        }

        payload = {
            "id":               nanoid(),
            "church_id":        CHURCH_ID,
            "full_name":        r["nom"],
            "nickname":         r["surnom"],
            "gender":           "male" if r["genre"] in ("H", "h") else "female",
            "phone":            phone_fmt,
            "location":         r["lieu"],
            "first_visit_date": r["prem_visite"],
            "spiritual_profile": spiritual_profile,
            "status":           "active",
            "is_undecided":     True,
        }
        payload = {k: v for k, v in payload.items() if v is not None and v != ""}
        resp = requests.post(url, headers={**HEADERS, "Prefer": "return=minimal"},
                             data=json.dumps(payload))
        if resp.status_code in (200, 201):
            inserted += 1
        else:
            errors.append(f"{r['nom']}: {resp.status_code} {resp.text[:80]}")

    print(f"\n✅ {inserted} âmes importées.")
    if errors:
        print(f"⚠️  {len(errors)} erreurs :")
        for e in errors:
            print(f"  {e}")

# ─── Main ───────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Import familles → Bergerie app")
    parser.add_argument("--import", dest="do_import", action="store_true",
                        help="Importer les âmes absentes après confirmation")
    parser.add_argument("--excel", action="store_true",
                        help="Exporter le rapport en Excel")
    args = parser.parse_args()

    print("Lecture du fichier Excel...")
    records = read_xlsx()
    print(f"  → {len(records)} lignes lues")

    print("Chargement des âmes depuis Supabase...")
    try:
        db_souls = fetch_souls()
        print(f"  → {len(db_souls)} âmes dans l'app")
    except Exception as e:
        print(f"  ❌ Erreur Supabase : {e}")
        sys.exit(1)

    results = compare(records, db_souls)
    totals  = print_report(results)

    if args.excel:
        export_excel(results)

    if args.do_import:
        do_import(results)
    elif totals.get("ABSENT", 0) > 0:
        print(f"💡 Pour importer les {totals['ABSENT']} âmes absentes :")
        print("   python scripts/import_familles.py --import")
        print()
        print("💡 Pour exporter le rapport en Excel :")
        print("   python scripts/import_familles.py --excel")

if __name__ == "__main__":
    main()
